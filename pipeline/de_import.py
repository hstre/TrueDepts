"""Import der amtlichen Rohdateien der Finanzagentur (Deutschland).

Liest die unveränderten Excel-Dateien aus data/raw/de/ und liefert normalisierte
Datensätze. Es werden keine Werte geschätzt; fehlende Angaben bleiben None und
werden in build.py gekennzeichnet.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl

from .bonds import add_business_days

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw" / "de"

INSTRUMENT_KIND = {
    "Bund": "fixed",
    "Bobl": "fixed",
    "Schatz": "fixed",
    "Green": "fixed",
    "Bubill": "zero",
    "ILB": "inflation_linked",
    "USD-Bond": "fixed_fx",
}

INSTRUMENT_LABEL = {
    "Bund": "Bundesanleihe",
    "Bobl": "Bundesobligation",
    "Schatz": "Bundesschatzanweisung",
    "Green": "Grünes Bundeswertpapier",
    "Bubill": "Unverzinsliche Schatzanweisung (Bubill)",
    "ILB": "Inflationsindexiertes Bundeswertpapier",
    "USD-Bond": "US-Dollar-Anleihe",
}

METHOD_LABEL = {
    "Auk": "Auktion",
    "M-A": "Multi-ISIN-Auktion",
    "Syn": "Syndikat",
    "EB": "Aufstockung in den Eigenbestand",
}


def _date(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    v = str(v).strip()
    m = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", v)
    if m:
        return dt.date(int(m[3]), int(m[2]), int(m[1]))
    return None


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    v = str(v).strip().replace(".", "").replace(",", ".")
    try:
        return float(v)
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Einzelaufstellung: Zinslaufbeginn je ISIN
# ---------------------------------------------------------------------------


def load_securities() -> dict[str, dict]:
    """ISIN -> {category, coupon, interest_start, maturity}. Aktuelle Liste ergänzt die Jahresliste."""
    out = {}
    for fname in ("einzelaufstellung_jahre_dt.xlsx", "einzelaufstellung_dt.xlsx"):
        ws = openpyxl.load_workbook(RAW / fname, read_only=True, data_only=True).worksheets[0]
        category = None
        dates = {}
        for r in ws.iter_rows(values_only=True):
            if not r:
                continue
            if r[0] == "WERTPAPIERART":
                dates = {j: _date(v) for j, v in enumerate(r) if j >= 5 and _date(v)}
                continue
            if r[0] and r[0] != "Zwischensumme":
                category = str(r[0]).strip()
            isin = str(r[1]).strip() if r[1] else ""
            if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}\d", isin):
                continue
            out.setdefault(isin, {
                "category": category,
                "coupon": round(_num(r[2]) / 100, 7) if _num(r[2]) is not None else None,
                "interest_start": _date(r[3]),
                "maturity": _date(r[4]),
                "source": fname,
                "outstanding": {},
            })
            for j, d in dates.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)):
                    out[isin]["outstanding"][d] = float(v) / 1e6
    return out


# ---------------------------------------------------------------------------
# Emissionshistorie
# ---------------------------------------------------------------------------


def settlement_date(auction_date, method, interest_start):
    """Valuta. Auktionen: zwei TARGET-Geschäftstage nach dem Auktionstag.

    Syndikate: Ist der Zinslaufbeginn bis zu 15 Tage nach dem Preisfeststellungstag, wird
    er als Valuta angenommen (Neuemission mit Zinslauf ab Valuta), sonst ebenfalls T+2.
    Die Datei enthält keine Valuta; das ist eine dokumentierte Annahme.
    """
    if method == "Syn" and interest_start and auction_date < interest_start <= auction_date + dt.timedelta(days=15):
        return interest_start, "Zinslaufbeginn (Syndikat)"
    return add_business_days(auction_date, 2), "T+2"


def load_auctions(securities=None) -> list[dict]:
    securities = securities if securities is not None else load_securities()
    ws = openpyxl.load_workbook(RAW / "emissionshistorie_dt.xlsx", read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    last_no = None
    for line, r in enumerate(rows, start=1):
        if not r or not r[2] or not isinstance(r[1], dt.datetime):
            continue
        isin = str(r[2]).strip()
        instrument = str(r[3]).strip()
        if r[0] not in (None, "", "  "):
            last_no = str(r[0]).strip()
        auction_date = r[1].date()
        method = str(r[9]).strip() if r[9] else None
        sec = securities.get(isin, {})
        interest_start = sec.get("interest_start")
        settle, settle_rule = settlement_date(auction_date, method, interest_start)
        issue_volume = _num(r[7]) or 0.0
        allotted = _num(r[13])
        retained = _num(r[17])
        if method == "EB":
            allotted = 0.0
            retained = retained if retained is not None else issue_volume
        rec = {
            "source_row": line,
            "auction_no": r[0] if r[0] not in (None, "", "  ") else None,
            "date": auction_date,
            "settle": settle,
            "settle_rule": settle_rule,
            "isin": isin,
            "instrument": instrument,
            "kind": INSTRUMENT_KIND.get(instrument, "unknown"),
            "coupon": _num(r[4]) if instrument != "Bubill" else 0.0,
            "maturity": _date(r[5]),
            "segment": str(r[6]).strip() if r[6] else None,
            "issue_volume": issue_volume,
            "new_issue": r[8] == "N",
            "method": method,
            "bids": _num(r[10]),
            "allotted": allotted if allotted is not None else 0.0,
            "retained": retained if retained is not None else 0.0,
            "price_low": _num(r[14]),
            "price_avg": _num(r[15]),
            "yield_avg": _num(r[16]),
            "bid_to_cover": _num(r[18]),
            "currency": "USD" if instrument == "USD-Bond" else "EUR",
            "interest_start": interest_start,
            "interest_start_source": sec.get("source"),
        }
        if rec["coupon"] is None and sec.get("coupon") is not None:
            rec["coupon"] = sec["coupon"]
        out.append(rec)
    # Neuemissionen ohne Eintrag in der Einzelaufstellung: Zinslaufbeginn = Valuta der ersten Emission
    first_settle = {}
    for rec in sorted(out, key=lambda x: x["date"]):
        first_settle.setdefault(rec["isin"], rec["settle"])
    for rec in out:
        if rec["interest_start"] is None and rec["kind"] != "zero":
            rec["interest_start"] = first_settle[rec["isin"]]
            rec["interest_start_source"] = "Annahme: Valuta der ersten Emission"
    return out


# ---------------------------------------------------------------------------
# Referenzindex / Index-Verhältniszahlen inflationsindexierter Bundeswertpapiere
# ---------------------------------------------------------------------------


def load_index_ratios() -> tuple[dict[str, dict[dt.date, float]], dict[str, dict]]:
    """ISIN -> {Datum: amtliche Index-Verhältniszahl}; ISIN -> Stammdaten aus den Kopfzeilen."""
    ratios: dict[str, dict[dt.date, float]] = {}
    meta: dict[str, dict] = {}
    for fname in ("archiv_referenzindex_bj2005_dt.xlsx", "archiv_referenzindex_bj2015_dt.xlsx",
                  "archiv_referenzindex_bj2025_dt.xlsx"):
        path = RAW / fname
        if not path.exists():
            continue
        rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(rows) if r and "Datum" in r)
        hdr = rows[hdr_i]
        date_col = hdr.index("Datum")
        cols = {}
        for j, h in enumerate(hdr):
            m = re.search(r"([A-Z]{2}[A-Z0-9]{9}\d)", str(h or ""))
            if m:
                cols[j] = m[1]
        for r in rows[hdr_i + 1: hdr_i + 8]:
            for j, isin in cols.items():
                v = str(r[j] or "")
                m = re.match(r"(Fälligkeit|Zinsen ab|Kupon|1\. Kupon|Basisindex): ([\d.,%]+)", v)
                if m:
                    meta.setdefault(isin, {})[f"{m[1]} ({fname[-11:-8]})" if m[1] == "Basisindex" else m[1]] = m[2]
        for r in rows[hdr_i + 1:]:
            d = r[date_col] if len(r) > date_col else None
            if not isinstance(d, dt.datetime):
                continue
            for j, isin in cols.items():
                v = r[j] if len(r) > j else None
                if isinstance(v, (int, float)):
                    ratios.setdefault(isin, {})[d.date()] = float(v)
    return ratios, meta


# ---------------------------------------------------------------------------
# Schuldenbericht: Jahreswerte
# ---------------------------------------------------------------------------

REPORT_SHEETS = {
    "gross_borrowing": "rpgBruttokreditaufnahme",
    "redemptions": "rpgTilgungen",
    "interest_cash": "rpgZinsen",
    "interest_total": "rpgZinsen Gesamt",
    "debt": "rpgSchuldenstand",
}
TOTAL_ROW = "Kredite Bundeshaushalt, Sondervermögen, Darlehensfinanzierung FMS & WSF"


def load_debt_report() -> dict:
    """Jahreswerte (Mio. €) aus dem Schuldenbericht. Stromgrößen sind im Bericht kumuliert
    seit Jahresbeginn; daher wird der Dezemberwert bzw. der letzte verfügbare Monat genommen.

    Vorzeichen werden so gedreht, dass Ausgaben/Tilgungen positiv sind.
    """
    wb = openpyxl.load_workbook(RAW / "schuldenbericht_dt.xlsx", read_only=True, data_only=True)
    years: dict[int, dict] = {}
    for key, sheet in REPORT_SHEETS.items():
        rows = list(wb[sheet].iter_rows(values_only=True))
        hdr = next(r for r in rows if r and r[0] == "Datum")
        top = next(r for r in rows if r and str(r[0] or "").startswith(TOTAL_ROW))
        by_year: dict[int, tuple[dt.date, float]] = {}
        for j in range(1, len(hdr)):
            d = _date(hdr[j])
            v = top[j] if j < len(top) else None
            if d is None or not isinstance(v, (int, float)):
                continue
            if d.year not in by_year or d > by_year[d.year][0]:
                by_year[d.year] = (d, float(v))
        for y, (d, v) in by_year.items():
            e = years.setdefault(y, {})
            val = v / 1e6
            if key in ("redemptions", "interest_cash", "interest_total"):
                val = -val
            e[key] = round(val, 1)
            e["as_of"] = max(e.get("as_of", ""), d.isoformat())
            e["complete_year"] = d.month == 12
    return years
